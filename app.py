
import streamlit as st
import sqlite3
import pandas as pd
import os
import re
import google.generativeai as genai
from dotenv import load_dotenv
from fpdf import FPDF
from datetime import datetime, date
from contextlib import contextmanager
from io import BytesIO
from docx import Document

# --- Config ---
load_dotenv()
api_key = os.getenv("GEMINI_API_KEY")
if api_key:
    genai.configure(api_key=api_key)
else:
    st.error("Chave da API do Gemini não encontrada no arquivo .env")

st.set_page_config(
    page_title="Gestão de Projetos - Mãe Natureza",
    page_icon="🌱",
    layout="wide"
)

# --- Constantes ---
DB_NAME = "ong_projetos.db"
UPLOAD_DIR = "uploads"
ORG_NAME = "Espaço Social Mãe Natureza"

STATUS_OPTIONS = ["Rascunho", "Em Preenchimento", "Em Revisão", "Submetido", "Aprovado", "Reprovado"]
STATUS_COLORS = {
    "Rascunho": "gray",
    "Em Preenchimento": "blue",
    "Em Revisão": "orange",
    "Submetido": "violet",
    "Aprovado": "green",
    "Reprovado": "red",
}
CATEGORIAS_ORCAMENTO = [
    "Recursos Humanos", "Material de Consumo", "Serviços de Terceiros",
    "Infraestrutura", "Comunicação", "Alimentação", "Transporte", "Geral"
]

if not os.path.exists(UPLOAD_DIR):
    os.makedirs(UPLOAD_DIR)


# --- Banco de Dados ---

@contextmanager
def get_db():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def migrate_db():
    """Adiciona colunas novas sem quebrar banco existente."""
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("PRAGMA table_info(projetos)")
        proj_cols = [row[1] for row in cursor.fetchall()]
        if "data_criacao" not in proj_cols:
            cursor.execute("ALTER TABLE projetos ADD COLUMN data_criacao TEXT")


migrate_db()


# --- Helpers de arquivo ---

def sanitize_filename(filename):
    filename = os.path.basename(filename)
    filename = re.sub(r"[^\w\s\-_\.]", "", filename).strip()
    return filename or "arquivo"


# --- Funções de dados ---

def get_projetos():
    with get_db() as conn:
        return pd.read_sql_query(
            """SELECT p.id, p.titulo, p.autor, p.status, p.data_criacao,
                      e.nome_edital, e.data_limite
               FROM projetos p
               LEFT JOIN editais e ON p.edital_id = e.id
               ORDER BY p.id DESC""",
            conn
        )


def get_projeto_by_id(projeto_id):
    with get_db() as conn:
        return conn.execute("SELECT * FROM projetos WHERE id = ?", (projeto_id,)).fetchone()


def get_editais():
    with get_db() as conn:
        return conn.execute("SELECT * FROM editais ORDER BY id DESC").fetchall()


def get_edital_by_id(edital_id):
    with get_db() as conn:
        return conn.execute("SELECT * FROM editais WHERE id = ?", (edital_id,)).fetchone()


def add_edital(nome, orgao, data_limite, pdf_path):
    with get_db() as conn:
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO editais (nome_edital, orgao_financiador, data_limite, arquivo_pdf_path) VALUES (?, ?, ?, ?)",
            (nome, orgao, data_limite, pdf_path)
        )
        return cur.lastrowid


def add_projeto(titulo, autor, edital_id):
    with get_db() as conn:
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO projetos (titulo, autor, status, edital_id, data_criacao) VALUES (?, ?, ?, ?, ?)",
            (titulo, autor, "Rascunho", edital_id, datetime.now().strftime("%d/%m/%Y %H:%M"))
        )
        return cur.lastrowid


def update_projeto_textos(projeto_id, justificativa, objetivos, publico_alvo):
    with get_db() as conn:
        conn.execute(
            "UPDATE projetos SET justificativa = ?, objetivos = ?, publico_alvo = ? WHERE id = ?",
            (justificativa, objetivos, publico_alvo, projeto_id)
        )


def update_projeto_status(projeto_id, novo_status):
    with get_db() as conn:
        conn.execute("UPDATE projetos SET status = ? WHERE id = ?", (novo_status, projeto_id))


def get_orcamento(projeto_id):
    with get_db() as conn:
        return pd.read_sql_query(
            "SELECT * FROM orcamentos WHERE projeto_id = ?", conn, params=(projeto_id,)
        )


def save_orcamento(projeto_id, df_orcamento):
    with get_db() as conn:
        conn.execute("DELETE FROM orcamentos WHERE projeto_id = ?", (projeto_id,))
        for _, row in df_orcamento.iterrows():
            item = str(row.get("item", "")).strip()
            if item:
                conn.execute(
                    "INSERT INTO orcamentos (projeto_id, item, valor, categoria) VALUES (?, ?, ?, ?)",
                    (projeto_id, item, max(0.0, float(row.get("valor", 0))), row.get("categoria", "Geral"))
                )


def add_anexo(projeto_id, nome, path):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO anexos (projeto_id, nome_arquivo, caminho_arquivo) VALUES (?, ?, ?)",
            (projeto_id, nome, path)
        )


def delete_projeto(projeto_id):
    with get_db() as conn:
        anexos = conn.execute(
            "SELECT caminho_arquivo FROM anexos WHERE projeto_id = ?", (projeto_id,)
        ).fetchall()
        for a in anexos:
            try:
                if os.path.exists(a["caminho_arquivo"]):
                    os.remove(a["caminho_arquivo"])
            except OSError:
                pass
        conn.execute("DELETE FROM orcamentos WHERE projeto_id = ?", (projeto_id,))
        conn.execute("DELETE FROM anexos WHERE projeto_id = ?", (projeto_id,))
        conn.execute("DELETE FROM projetos WHERE id = ?", (projeto_id,))


# --- Gemini ---

def get_gemini_model(system_instruction):
    return genai.GenerativeModel("gemini-2.0-flash", system_instruction=system_instruction)


def get_gemini_file(pdf_path):
    """Faz upload do PDF para o Gemini, com cache por caminho na session."""
    cache_key = f"gemini_file_{pdf_path}"
    if cache_key not in st.session_state:
        st.session_state[cache_key] = genai.upload_file(pdf_path, mime_type="application/pdf")
    return st.session_state[cache_key]


# --- UI Helpers ---

def status_badge(status):
    color = STATUS_COLORS.get(status, "gray")
    return f":{color}[**{status}**]"


def mostrar_prazo(data_limite_str):
    """Exibe alerta colorido baseado no prazo do edital."""
    if not data_limite_str:
        return
    try:
        prazo = datetime.strptime(data_limite_str, "%Y-%m-%d").date()
        dias = (prazo - date.today()).days
        if dias < 0:
            st.error(f"⛔ Prazo encerrado em {prazo.strftime('%d/%m/%Y')}")
        elif dias <= 7:
            st.warning(f"⏰ Prazo se encerra em {prazo.strftime('%d/%m/%Y')} — **{dias} dias restantes!**")
        elif dias <= 30:
            st.info(f"📅 Prazo: {prazo.strftime('%d/%m/%Y')} ({dias} dias restantes)")
        else:
            st.caption(f"📅 Prazo de submissão: {prazo.strftime('%d/%m/%Y')}")
    except ValueError:
        pass


# --- Páginas ---

def dashboard():
    st.title("📊 Painel de Projetos")

    df = get_projetos()

    total = len(df)
    aprovados = len(df[df["status"] == "Aprovado"])
    reprovados = len(df[df["status"] == "Reprovado"])
    ativos = len(df[~df["status"].isin(["Aprovado", "Reprovado"])])

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total", total)
    c2.metric("Ativos", ativos)
    c3.metric("Aprovados", aprovados)
    c4.metric("Reprovados", reprovados)

    st.markdown("---")
    st.subheader("Projetos")

    if df.empty:
        st.info("Nenhum projeto cadastrado. Use **Novo Projeto** para começar.")
        return

    # Filtros
    col_busca, col_filtro = st.columns([3, 1])
    with col_busca:
        busca = st.text_input("🔍 Buscar por título", placeholder="Digite para filtrar...")
    with col_filtro:
        filtro_status = st.selectbox("Status", ["Todos"] + STATUS_OPTIONS)

    df_vis = df.copy()
    if busca:
        df_vis = df_vis[df_vis["titulo"].str.contains(busca, case=False, na=False)]
    if filtro_status != "Todos":
        df_vis = df_vis[df_vis["status"] == filtro_status]

    if df_vis.empty:
        st.info("Nenhum projeto encontrado para os filtros selecionados.")
    else:
        for _, row in df_vis.iterrows():
            with st.container(border=True):
                col_a, col_b, col_c, col_d, col_e, col_f = st.columns([3, 2, 2, 1.5, 0.6, 0.6])
                with col_a:
                    st.write(f"**{row['titulo']}**")
                    if row.get("data_criacao"):
                        st.caption(f"Criado em: {row['data_criacao']}")
                with col_b:
                    st.write(row.get("nome_edital") or "—")
                    if row.get("data_limite"):
                        try:
                            prazo = datetime.strptime(row["data_limite"], "%Y-%m-%d").date()
                            dias = (prazo - date.today()).days
                            if dias < 0:
                                st.caption(f"⛔ Prazo encerrado")
                            elif dias <= 7:
                                st.caption(f"⏰ {dias} dias restantes")
                        except ValueError:
                            pass
                with col_c:
                    st.write(row.get("autor") or "—")
                with col_d:
                    st.markdown(status_badge(row["status"]))
                with col_e:
                    if st.button("✏️", key=f"edit_{row['id']}", help="Editar"):
                        st.session_state["navigate_to"] = "Revisão e Orçamento"
                        st.session_state["editing_project_id"] = row["id"]
                        st.rerun()
                with col_f:
                    if st.button("🗑️", key=f"del_{row['id']}", help="Excluir"):
                        st.session_state["confirmar_delete_id"] = row["id"]
                        st.session_state["confirmar_delete_titulo"] = row["titulo"]
                        st.rerun()

    # Confirmação de exclusão
    if st.session_state.get("confirmar_delete_id"):
        pid = st.session_state["confirmar_delete_id"]
        titulo = st.session_state.get("confirmar_delete_titulo", "")
        st.markdown("---")
        with st.container(border=True):
            st.error(f"⚠️ Excluir o projeto **\"{titulo}\"**? Esta ação é irreversível.")
            col_sim, col_nao, _ = st.columns([1, 1, 5])
            with col_sim:
                if st.button("Sim, excluir", type="primary"):
                    delete_projeto(pid)
                    st.session_state.pop("confirmar_delete_id", None)
                    st.session_state.pop("confirmar_delete_titulo", None)
                    st.rerun()
            with col_nao:
                if st.button("Cancelar"):
                    st.session_state.pop("confirmar_delete_id", None)
                    st.session_state.pop("confirmar_delete_titulo", None)
                    st.rerun()


def setup_projeto():
    st.title("📝 Novo Projeto")

    # 1. Seleção do edital
    st.subheader("1. Seleção do Edital")

    editais = get_editais()
    opcoes = {e["id"]: f"{e['nome_edital']} ({e['orgao_financiador'] or 'S/ Órgão'})" for e in editais}
    opcoes[0] = "➕ Cadastrar Novo Edital"

    if "selected_edital_id" not in st.session_state:
        st.session_state.selected_edital_id = None

    edital_selection = st.selectbox(
        "Selecione o Edital Base:",
        options=list(opcoes.keys()),
        format_func=lambda x: opcoes[x],
        index=0 if not editais else None
    )

    if edital_selection == 0:
        with st.expander("Cadastro de Novo Edital", expanded=True):
            with st.form("novo_edital_form"):
                novo_nome = st.text_input("Nome do Edital *")
                novo_orgao = st.text_input("Órgão Financiador")
                nova_data = st.date_input("Data Limite de Submissão", value=None)
                novo_pdf = st.file_uploader("PDF do Edital *", type="pdf")

                if st.form_submit_button("Salvar Edital", type="primary"):
                    if novo_nome and novo_pdf:
                        safe = sanitize_filename(novo_pdf.name)
                        path = os.path.join(UPLOAD_DIR, safe)
                        with open(path, "wb") as f:
                            f.write(novo_pdf.getbuffer())
                        data_str = nova_data.strftime("%Y-%m-%d") if nova_data else None
                        new_id = add_edital(novo_nome, novo_orgao, data_str, path)
                        st.session_state.selected_edital_id = new_id
                        st.success(f"Edital '{novo_nome}' cadastrado!")
                        st.rerun()
                    else:
                        st.error("Nome e PDF são obrigatórios.")
    else:
        st.session_state.selected_edital_id = edital_selection

    # Mostrar prazo do edital selecionado
    if st.session_state.get("selected_edital_id") and edital_selection != 0:
        ed = get_edital_by_id(st.session_state.selected_edital_id)
        if ed:
            mostrar_prazo(ed["data_limite"])

    st.markdown("---")

    # 2. Dados do projeto
    st.subheader("2. Dados do Projeto")

    if st.session_state.get("selected_edital_id") and edital_selection != 0:
        sel_edital = get_edital_by_id(st.session_state.selected_edital_id)
        st.info(f"Vinculado ao Edital: **{sel_edital['nome_edital']}**")

        with st.form("novo_projeto_form"):
            titulo = st.text_input("Nome do Projeto *")
            autor = st.text_input("Responsável / Autor *")
            anexos = st.file_uploader(
                "Anexos Extras (Planilhas, Documentos, Imagens)",
                accept_multiple_files=True
            )

            if st.form_submit_button("Criar Projeto", type="primary"):
                if titulo and autor:
                    projeto_id = add_projeto(titulo, autor, st.session_state.selected_edital_id)
                    if anexos:
                        for arq in anexos:
                            safe = sanitize_filename(arq.name)
                            path = os.path.join(UPLOAD_DIR, f"{projeto_id}_{safe}")
                            with open(path, "wb") as f:
                                f.write(arq.getbuffer())
                            add_anexo(projeto_id, arq.name, path)
                    st.success(f"✅ Projeto **'{titulo}'** criado! Agora use o **Assistente IA** para elaborar a proposta.")
                else:
                    st.error("Nome do projeto e responsável são obrigatórios.")
    else:
        st.info("Selecione ou cadastre um edital acima para continuar.")


def assistente_ia():
    st.title("🤖 Chat com Rafa Ia")

    projetos = get_projetos()
    if projetos.empty:
        st.warning("Nenhum projeto encontrado. Crie um projeto primeiro.")
        return

    projeto_id = st.selectbox(
        "Selecione o Projeto",
        options=projetos["id"],
        format_func=lambda x: projetos[projetos["id"] == x]["titulo"].values[0]
    )
    projeto = get_projeto_by_id(projeto_id)
    if not projeto:
        return

    edital = get_edital_by_id(projeto["edital_id"])

    col_info, col_prazo = st.columns([3, 1])
    with col_info:
        st.info(f"Trabalhando em: **{projeto['titulo']}** | Edital: **{edital['nome_edital']}**")
    with col_prazo:
        mostrar_prazo(edital["data_limite"] if edital else None)

    pdf_path = edital["arquivo_pdf_path"] if edital else None
    if not pdf_path:
        st.error("O edital não possui PDF vinculado. Cadastre o PDF na seção Novo Projeto.")
        return
    if not os.path.exists(pdf_path):
        st.error(f"Arquivo PDF não encontrado: {pdf_path}")
        return

    st.markdown("---")

    # Inicializar chat apenas quando muda de projeto
    if "chat_session" not in st.session_state or st.session_state.get("current_project_id") != projeto_id:
        st.session_state.current_project_id = projeto_id
        st.session_state.messages = []

        with st.spinner("Rafa Ia está lendo o Edital..."):
            try:
                gemini_file = get_gemini_file(pdf_path)

                system_instruction = f"""
                Você é a **Rafa Ia**, a inteligência artificial do {ORG_NAME}.
                Seu objetivo é entrevistar o usuário para coletar informações e montar um projeto social/cultural.
                Use o EDITAL fornecido como referência principal — verifique os requisitos, valores, prazos e critérios.

                Conduza a entrevista em fases:
                1. Ideia central (Justificativa) — contexto, problema e relevância
                2. Objetivos — geral e específicos, alinhados ao edital
                3. Público-alvo — quem será beneficiado e como

                Seja simpática e use emojis moderadamente. Faça uma pergunta de cada vez.
                Ao final, apresente um resumo completo para o usuário validar.
                """

                model = get_gemini_model(system_instruction)
                st.session_state.chat = model.start_chat(history=[
                    {"role": "user", "parts": [gemini_file, "Analise este edital e comece a entrevista."]}
                ])
                response = st.session_state.chat.send_message("Olá! Pode se apresentar e começar?")
                st.session_state.messages.append({"role": "assistant", "content": response.text})
                st.session_state.chat_session = True
            except Exception as e:
                st.error(f"Erro ao conectar com Rafa Ia: {e}")
                return

    # Histórico do chat
    for msg in st.session_state.messages:
        avatar = "🤖" if msg["role"] == "assistant" else "👤"
        with st.chat_message(msg["role"], avatar=avatar):
            st.markdown(msg["content"])

    if prompt := st.chat_input("Fale com a Rafa Ia..."):
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user", avatar="👤"):
            st.markdown(prompt)

        if "chat" in st.session_state:
            with st.chat_message("assistant", avatar="🤖"):
                with st.spinner("Rafa Ia está digitando..."):
                    try:
                        response = st.session_state.chat.send_message(prompt)
                        st.markdown(response.text)
                        st.session_state.messages.append({"role": "assistant", "content": response.text})
                    except Exception as e:
                        st.error(f"Erro na resposta: {e}")


def revisao_orcamento():
    st.title("💰 Revisão e Orçamento")

    projetos = get_projetos()
    if projetos.empty:
        st.warning("Nenhum projeto encontrado.")
        return

    list_ids = projetos["id"].tolist()
    default_index = 0
    editing_id = st.session_state.get("editing_project_id")
    if editing_id in list_ids:
        default_index = list_ids.index(editing_id)

    projeto_id = st.selectbox(
        "Selecione o Projeto",
        options=projetos["id"],
        format_func=lambda x: projetos[projetos["id"] == x]["titulo"].values[0],
        index=default_index,
        key="rev_proj_sel"
    )
    st.session_state["editing_project_id"] = projeto_id

    projeto = get_projeto_by_id(projeto_id)
    edital = get_edital_by_id(projeto["edital_id"]) if projeto else None

    # Barra de status + prazo
    col_s1, col_s2, col_s3 = st.columns([2, 2, 2])
    with col_s1:
        status_atual = projeto["status"] if projeto else "Rascunho"
        st.markdown(f"**Status:** {status_badge(status_atual)}")
        if projeto and projeto.get("autor"):
            st.caption(f"Responsável: {projeto['autor']}")
    with col_s2:
        novo_status = st.selectbox(
            "Alterar status",
            STATUS_OPTIONS,
            index=STATUS_OPTIONS.index(status_atual) if status_atual in STATUS_OPTIONS else 0,
            key="status_sel"
        )
        if st.button("Atualizar Status"):
            update_projeto_status(projeto_id, novo_status)
            st.success(f"Status atualizado para **{novo_status}**")
            st.rerun()
    with col_s3:
        mostrar_prazo(edital["data_limite"] if edital else None)

    st.markdown("---")

    tab1, tab2, tab3 = st.tabs(["📝 Textos do Projeto", "💸 Orçamento", "📤 Exportar"])

    with tab1:
        st.subheader("Refinamento da Proposta")

        with st.form("textos_form"):
            justificativa = st.text_area(
                "Justificativa",
                value=projeto["justificativa"] or "",
                height=200,
                help="Explique o contexto, o problema e a relevância do projeto."
            )
            objetivos = st.text_area(
                "Objetivos",
                value=projeto["objetivos"] or "",
                height=200,
                help="Liste o objetivo geral e os objetivos específicos."
            )
            publico = st.text_area(
                "Público-Alvo",
                value=projeto["publico_alvo"] or "",
                height=100,
                help="Descreva quem será beneficiado e como."
            )
            if st.form_submit_button("💾 Salvar Textos", type="primary"):
                update_projeto_textos(projeto_id, justificativa, objetivos, publico)
                st.success("Textos salvos com sucesso!")

        st.markdown("---")
        st.subheader("✨ Análise da Rafa Ia")
        if st.button("Analisar Proposta"):
            with st.spinner("Rafa Ia está analisando..."):
                texto = (
                    f"Projeto: {projeto['titulo']}\n\n"
                    f"Justificativa: {justificativa}\n\n"
                    f"Objetivos: {objetivos}\n\n"
                    f"Público-alvo: {publico}"
                )
                prompt = f"""
                Você é uma avaliadora experiente de projetos sociais e culturais.
                Analise o rascunho abaixo e dê feedback construtivo.
                Foque em: Clareza, Impacto Social, Coerência e pontos de melhoria práticos.

                {texto}

                Apresente: pontos fortes, pontos de melhoria e sugestões concretas.
                """
                try:
                    model = get_gemini_model("Avaliadora experiente de projetos sociais.")
                    response = model.generate_content(prompt)
                    st.info(response.text)
                except Exception as e:
                    st.error(f"Erro ao analisar: {e}")

    with tab2:
        st.subheader("Planilha Orçamentária")

        df_orc = get_orcamento(projeto_id)
        if df_orc.empty:
            df_orc = pd.DataFrame([{"item": "", "valor": 0.0, "categoria": "Geral"}])
        else:
            df_orc = df_orc[["item", "valor", "categoria"]]

        edited_df = st.data_editor(
            df_orc,
            num_rows="dynamic",
            use_container_width=True,
            column_config={
                "item": st.column_config.TextColumn("Item / Descrição", required=True),
                "valor": st.column_config.NumberColumn(
                    "Valor (R$)", min_value=0.0, format="R$ %.2f"
                ),
                "categoria": st.column_config.SelectboxColumn(
                    "Categoria", options=CATEGORIAS_ORCAMENTO
                ),
            }
        )

        linhas_validas = edited_df[edited_df["item"].astype(str).str.strip() != ""]
        total = linhas_validas["valor"].sum()
        st.metric("Valor Total do Projeto", f"R$ {total:,.2f}")

        if st.button("💾 Salvar Orçamento", type="primary"):
            save_orcamento(projeto_id, edited_df)
            st.success("Orçamento salvo com sucesso!")

    with tab3:
        st.subheader("Exportar Proposta Final")
        col_pdf, col_docx = st.columns(2)

        with col_pdf:
            st.write("**PDF**")
            if st.button("Gerar PDF", use_container_width=True):
                pdf_bytes = _gerar_pdf(projeto_id)
                if pdf_bytes:
                    nome_arquivo = projeto["titulo"].replace(" ", "_") if projeto else "proposta"
                    st.download_button(
                        label="⬇️ Baixar PDF",
                        data=pdf_bytes,
                        file_name=f"proposta_{nome_arquivo}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

        with col_docx:
            st.write("**Word (DOCX)**")
            if st.button("Gerar DOCX", use_container_width=True):
                docx_bytes = _gerar_docx(projeto_id)
                if docx_bytes:
                    nome_arquivo = projeto["titulo"].replace(" ", "_") if projeto else "proposta"
                    st.download_button(
                        label="⬇️ Baixar DOCX",
                        data=docx_bytes,
                        file_name=f"proposta_{nome_arquivo}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )


# --- Geração de arquivos ---

def _limpar_texto(texto):
    """Converte para latin-1 de forma segura para o fpdf 1.x."""
    if not texto:
        return "Não preenchido."
    return str(texto).encode("latin-1", "replace").decode("latin-1")


def _gerar_pdf(projeto_id):
    projeto = get_projeto_by_id(projeto_id)
    edital = get_edital_by_id(projeto["edital_id"]) if projeto else None
    orcamento = get_orcamento(projeto_id)

    pdf = FPDF()
    pdf.add_page()

    # Cabeçalho
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, _limpar_texto(ORG_NAME), 0, 1, "C")
    pdf.set_font("Arial", "", 10)
    if edital:
        pdf.cell(0, 6, _limpar_texto(f"Edital: {edital['nome_edital']}"), 0, 1, "C")
        if edital["orgao_financiador"]:
            pdf.cell(0, 6, _limpar_texto(f"Orgao: {edital['orgao_financiador']}"), 0, 1, "C")
        if edital["data_limite"]:
            pdf.cell(0, 6, _limpar_texto(f"Prazo de Submissao: {edital['data_limite']}"), 0, 1, "C")
    pdf.ln(6)

    # Título
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, _limpar_texto(f"Projeto: {projeto['titulo']}"), 0, 1, "L")
    pdf.set_font("Arial", "", 11)
    if projeto["autor"]:
        pdf.cell(0, 7, _limpar_texto(f"Responsavel: {projeto['autor']}"), 0, 1, "L")
    pdf.cell(0, 7, _limpar_texto(f"Status: {projeto['status']}"), 0, 1, "L")
    if projeto.get("data_criacao"):
        pdf.cell(0, 7, _limpar_texto(f"Criado em: {projeto['data_criacao']}"), 0, 1, "L")
    pdf.ln(4)

    def add_section(title, content):
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, _limpar_texto(title), 0, 1, "L")
        pdf.set_font("Arial", "", 11)
        pdf.multi_cell(0, 7, _limpar_texto(content))
        pdf.ln(4)

    add_section("Justificativa", projeto["justificativa"])
    add_section("Objetivos", projeto["objetivos"])
    add_section("Publico-Alvo", projeto["publico_alvo"])

    # Orçamento
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Orcamento Detalhado", 0, 1, "L")
    pdf.ln(2)
    pdf.set_font("Arial", "B", 10)
    pdf.cell(90, 7, "Item", 1)
    pdf.cell(40, 7, "Categoria", 1)
    pdf.cell(40, 7, "Valor (R$)", 1)
    pdf.ln()

    pdf.set_font("Arial", "", 10)
    total_pdf = 0.0
    if not orcamento.empty:
        for _, row in orcamento.iterrows():
            pdf.cell(90, 7, _limpar_texto(str(row["item"])[:50]), 1)
            pdf.cell(40, 7, _limpar_texto(str(row["categoria"])[:25]), 1)
            pdf.cell(40, 7, f"{row['valor']:.2f}", 1)
            pdf.ln()
            total_pdf += row["valor"]

    pdf.ln(4)
    pdf.set_font("Arial", "B", 11)
    pdf.cell(0, 10, f"Valor Total: R$ {total_pdf:,.2f}", 0, 1, "R")

    try:
        return bytes(pdf.output())
    except Exception:
        return pdf.output(dest="S").encode("latin-1")


def _gerar_docx(projeto_id):
    projeto = get_projeto_by_id(projeto_id)
    edital = get_edital_by_id(projeto["edital_id"]) if projeto else None
    orcamento = get_orcamento(projeto_id)

    doc = Document()
    doc.add_heading(ORG_NAME, 0)
    if edital:
        doc.add_paragraph(f"Edital: {edital['nome_edital']}")
        if edital["orgao_financiador"]:
            doc.add_paragraph(f"Órgão: {edital['orgao_financiador']}")
        if edital["data_limite"]:
            doc.add_paragraph(f"Prazo de Submissão: {edital['data_limite']}")

    doc.add_heading(projeto["titulo"], level=1)
    if projeto["autor"]:
        doc.add_paragraph(f"Responsável: {projeto['autor']}")
    doc.add_paragraph(f"Status: {projeto['status']}")

    for titulo, campo in [
        ("Justificativa", projeto["justificativa"]),
        ("Objetivos", projeto["objetivos"]),
        ("Público-Alvo", projeto["publico_alvo"]),
    ]:
        doc.add_heading(titulo, level=2)
        doc.add_paragraph(campo or "Não preenchido.")

    doc.add_heading("Orçamento", level=2)
    if not orcamento.empty:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Item", "Valor (R$)", "Categoria"
        for _, row in orcamento.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["item"])
            cells[1].text = f"R$ {row['valor']:.2f}"
            cells[2].text = str(row["categoria"])
        doc.add_paragraph(f"\nValor Total: R$ {orcamento['valor'].sum():,.2f}")
    else:
        doc.add_paragraph("Nenhum item de orçamento cadastrado.")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# --- Navegação Principal ---

def main():
    st.sidebar.image("https://placehold.co/200x60?text=MAE+NATUREZA", width=200)
    st.sidebar.title("Gestão de Projetos")

    if "page" not in st.session_state:
        st.session_state.page = "Dashboard"

    if "navigate_to" in st.session_state:
        st.session_state.page = st.session_state.pop("navigate_to")
        st.rerun()

    page = st.sidebar.radio(
        "Navegação",
        ["Dashboard", "Novo Projeto", "Assistente IA", "Revisão e Orçamento"],
        key="page"
    )

    if page == "Dashboard":
        dashboard()
    elif page == "Novo Projeto":
        setup_projeto()
    elif page == "Assistente IA":
        assistente_ia()
    elif page == "Revisão e Orçamento":
        revisao_orcamento()


if __name__ == "__main__":
    main()
